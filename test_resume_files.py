# test_resume_files.py - checks for reading resume files (all made up in memory)
from io import BytesIO

import pytest
from docx import Document

from resume_files import read_resume


def test_markdown_resume():
    # b"..." is bytes: what a file really holds, before it's turned into text
    text = read_resume("resume.md", b"# Alex\n- Ran on-call\n")
    assert text == "# Alex\n- Ran on-call"


def test_word_resume():
    document = Document()  # a new, empty Word document
    document.add_paragraph("Alex Rivera")
    document.add_paragraph("Led incident management for online banking")
    buffer = BytesIO()  # an in-memory "file"
    document.save(buffer)
    text = read_resume("Resume.DOCX", buffer.getvalue())
    assert text == "Alex Rivera\nLed incident management for online banking"


def test_other_files_get_a_clear_message():
    with pytest.raises(ValueError, match="Use PDF"):
        read_resume("resume.pages", b"...")